# -*- coding: utf-8 -*-
"""
将 ods_ddl.sql 中的 CREATE TABLE 解析为业务友好的 Excel 数据字典。
优化点：
  1. 首页"目录"汇总所有表（表名、业务说明、字段数）
  2. 列名业务化：是否必填、业务说明等
  3. 数据类型增加"业务含义"列（如 varchar(64) → 字符串）
  4. 样式：表头蓝底白字、边框、自动换行、冻结窗格、主键行高亮
"""
import re
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

SQL_FILE = Path(r"e:\project\req\ShangLian\ods_ddl.sql")
XLSX_FILE = Path(r"e:\project\req\ShangLian\ods_ddl_业务版.xlsx")

text = SQL_FILE.read_text(encoding="utf-8")

# 数据类型 → 业务含义映射
TYPE_DESC = {
    "varchar": "字符串",
    "char": "字符串",
    "text": "长文本",
    "int": "整数",
    "bigint": "长整数",
    "tinyint": "整数",
    "smallint": "整数",
    "decimal": "小数",
    "date": "日期",
    "datetime": "日期时间",
    "timestamp": "时间戳",
    "time": "时间",
}

def get_type_desc(data_type):
    """根据数据类型返回业务含义"""
    base = data_type.split("(")[0].strip().lower()
    desc = TYPE_DESC.get(base, base)
    # 提取长度/精度
    m = re.search(r"\(([^)]*)\)", data_type)
    if m:
        params = m.group(1).replace(" ", "")
        if base in ("varchar", "char"):
            desc += f"（最长{params}字符）"
        elif base == "decimal":
            parts = params.split(",")
            if len(parts) == 2:
                desc += f"（总{parts[0]}位，小数{parts[1]}位）"
        elif base in ("int", "bigint", "tinyint"):
            pass  # 整数不需要显示长度
    return desc

# 按 CREATE TABLE 切分，每个块一张表
blocks = re.split(r"CREATE TABLE\s+`", text, flags=re.IGNORECASE)
tables = []
for block in blocks[1:]:
    tbl_name = block.split("`", 1)[0]
    tail_m = re.search(r"DISTRIBUTE\s+BY.*$", block, re.DOTALL | re.IGNORECASE)
    tail = tail_m.group(0) if tail_m else ""
    tbl_comment_m = re.search(r"COMMENT\s*=?\s*'((?:[^'\\]|\\.)*)'", tail, re.IGNORECASE)
    tbl_comment = tbl_comment_m.group(1) if tbl_comment_m else ""

    inner_m = re.search(r"\((.*)\)\s*DISTRIBUTE\s+BY", block, re.DOTALL | re.IGNORECASE)
    if not inner_m:
        inner_m = re.search(r"\((.*)\)\s*;?\s*$", block, re.DOTALL | re.IGNORECASE)
    if not inner_m:
        continue
    inner = inner_m.group(1)

    cols = []
    depth = 0
    cur = ""
    for ch in inner:
        if ch == "(":
            depth += 1
            cur += ch
        elif ch == ")":
            depth -= 1
            cur += ch
        elif ch == "," and depth == 0:
            cols.append(cur.strip())
            cur = ""
        else:
            cur += ch
    if cur.strip():
        cols.append(cur.strip())

    pk_fields = set()
    pk_m = re.search(r"PRIMARY\s+KEY\s*\(([^)]*)\)", inner, re.IGNORECASE)
    if pk_m:
        pk_fields = set(x.strip().strip("`") for x in pk_m.group(1).split(","))

    rows = []
    idx = 0
    for col in cols:
        col_stripped = col.strip()
        if re.match(r"^(KEY|INDEX|UNIQUE|FULLTEXT|SPATIAL|PRIMARY\s+KEY|CONSTRAINT|FOREIGN\s+KEY)\b",
                    col_stripped, re.IGNORECASE):
            continue
        fm = re.match(r"`([^`]+)`\s+(.+)", col_stripped, re.DOTALL)
        if not fm:
            continue
        field_name = fm.group(1)
        rest = fm.group(2).rstrip(",").strip()

        type_m = re.match(r"([A-Za-z]+(?:\s*\([^)]*\))?)", rest)
        data_type = type_m.group(1) if type_m else ""

        is_pk = "是" if field_name in pk_fields else "否"
        is_not_null = "是" if re.search(r"\bNOT\s+NULL\b", rest, re.IGNORECASE) else "否"

        default = ""
        default_m = re.search(r"\bDEFAULT\s+('([^']*)'|[-+]?\d+\.?\d*|CURRENT_TIMESTAMP|NULL|TRUE|FALSE)",
                               rest, re.IGNORECASE)
        if default_m:
            default = default_m.group(1)

        comment = ""
        comment_m = re.search(r"COMMENT\s+'((?:[^'\\]|\\.)*)'", rest, re.IGNORECASE)
        if comment_m:
            comment = comment_m.group(1).replace("\\'", "'")

        idx += 1
        rows.append({
            "序号": idx,
            "字段名称": field_name,
            "业务说明": comment,
            "数据类型": data_type,
            "类型说明": get_type_desc(data_type),
            "是否必填": is_not_null,
            "是否主键": is_pk,
            "默认值": default,
        })

    tables.append((tbl_name, tbl_comment, rows))

# ============ 写 Excel ============
with pd.ExcelWriter(XLSX_FILE, engine="openpyxl") as writer:
    # --- 目录 sheet ---
    catalog_rows = []
    for i, (tbl_name, tbl_comment, rows) in enumerate(tables, 1):
        catalog_rows.append({
            "序号": i,
            "表名": tbl_name,
            "业务说明": tbl_comment,
            "字段数量": len(rows),
        })
    pd.DataFrame(catalog_rows).to_excel(writer, sheet_name="目录", index=False)

    # --- 各表 sheet ---
    for tbl_name, tbl_comment, rows in tables:
        df = pd.DataFrame(rows)
        sheet = tbl_name[:31]
        df.to_excel(writer, sheet_name=sheet, index=False)

# ============ 统一样式美化 ============
wb = load_workbook(XLSX_FILE)

# 样式定义
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
title_font = Font(name="微软雅黑", size=12, bold=True, color="1F4E79")
cell_font = Font(name="微软雅黑", size=10)
pk_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # 主键行浅黄
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

def apply_table_style(ws, n_cols, has_pk_col=True):
    """对单张表 sheet 应用样式"""
    # 表头行（第1行）
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 数据行
    pk_col_idx = None
    for col in range(1, n_cols + 1):
        if ws.cell(row=1, column=col).value == "是否主键":
            pk_col_idx = col
            break

    for row in range(2, ws.max_row + 1):
        is_pk_row = False
        if pk_col_idx:
            is_pk_row = (ws.cell(row=row, column=pk_col_idx).value == "是")

        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col in (1, 5, 6, 7):  # 序号、类型说明、是否必填、是否主键 居中
                cell.alignment = center_align
            else:
                cell.alignment = left_align
            if is_pk_row:
                cell.fill = pk_fill

    # 冻结表头
    ws.freeze_panes = "A2"
    # 自动筛选
    ws.auto_filter.ref = ws.dimensions

def set_col_widths(ws, widths):
    """设置列宽"""
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

# 美化目录 sheet
ws_cat = wb["目录"]
n_cat_cols = ws_cat.max_column
for col in range(1, n_cat_cols + 1):
    cell = ws_cat.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border
for row in range(2, ws_cat.max_row + 1):
    for col in range(1, n_cat_cols + 1):
        cell = ws_cat.cell(row=row, column=col)
        cell.font = cell_font
        cell.border = thin_border
        cell.alignment = center_align if col in (1, 4) else left_align
ws_cat.freeze_panes = "A2"
ws_cat.auto_filter.ref = ws_cat.dimensions
set_col_widths(ws_cat, [8, 32, 40, 12])

# 美化各表 sheet
table_widths = [8, 30, 50, 18, 24, 12, 12, 18]
for tbl_name, tbl_comment, rows in tables:
    sheet = tbl_name[:31]
    ws = wb[sheet]
    apply_table_style(ws, ws.max_column)
    set_col_widths(ws, table_widths)

wb.save(XLSX_FILE)

print(f"已生成: {XLSX_FILE}")
print(f"共 {len(tables)} 张表:")
for name, comment, rows in tables:
    print(f"  - {name}（{comment}）：{len(rows)} 个字段")
